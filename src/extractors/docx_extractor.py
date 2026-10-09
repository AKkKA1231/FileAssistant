"""Microsoft Word DOCX text extraction."""

from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from contracts import ErrorCode
from extractors.base import ExtractionError, ExtractionResult


def extract_docx(path: Path) -> ExtractionResult:
    """Extract paragraphs and table rows from a DOCX document."""

    try:
        document = Document(path)
    except (BadZipFile, KeyError, PackageNotFoundError, ValueError) as exc:
        raise ExtractionError(
            ErrorCode.PARSE_ERROR,
            "The DOCX file is corrupt or could not be parsed.",
            details={"path": str(path)},
        ) from exc

    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    table_rows: list[str] = []
    for table in document.tables:
        for row in table.rows:
            table_rows.append("\t".join(cell.text for cell in row.cells))

    sections = paragraphs + table_rows
    content = "\n".join(sections)
    warnings: list[str] = []
    if not content.strip():
        warnings.append("No extractable text was found in the DOCX file.")

    return ExtractionResult(
        content=content,
        details={
            "format": "docx",
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
        },
        warnings=tuple(warnings),
    )


__all__ = ["extract_docx"]

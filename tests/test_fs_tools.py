"""Unit and integration tests for the core file-system tools."""

from __future__ import annotations

from pathlib import Path
from typing import Any, cast

from docx import Document

from fs_tools import list_files, read_file, search_in_file, write_file


def _data(response: dict[str, Any]) -> dict[str, Any]:
    assert response["success"] is True
    return cast(dict[str, Any], response["data"])


def _error_code(response: dict[str, Any]) -> str:
    assert response["success"] is False
    error = cast(dict[str, Any], response["error"])
    return cast(str, error["code"])


def _write_minimal_text_pdf(path: Path, text: str) -> None:
    escaped_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 18 Tf 72 720 Td ({escaped_text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{object_number} 0 obj\n".encode())
        pdf.extend(body)
        pdf.extend(b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    path.write_bytes(pdf)


def test_read_utf8_txt_returns_content_and_metadata(tmp_path: Path) -> None:
    resume = tmp_path / "resume.txt"
    resume.write_text("Python developer", encoding="utf-8")

    response = read_file("resume.txt", workspace_root=tmp_path)
    data = _data(response)

    assert data["content"] == "Python developer"
    assert data["metadata"]["name"] == "resume.txt"
    assert data["extraction"] == {"format": "txt", "encoding": "utf-8"}


def test_read_txt_uses_cp1252_fallback(tmp_path: Path) -> None:
    (tmp_path / "resume.txt").write_bytes("André".encode("cp1252"))

    response = read_file("resume.txt", workspace_root=tmp_path)
    data = _data(response)

    assert data["content"] == "André"
    assert data["extraction"]["encoding"] == "cp1252"
    assert response["warnings"]


def test_read_txt_rejects_binary_data(tmp_path: Path) -> None:
    (tmp_path / "binary.txt").write_bytes(b"text\x00binary")

    response = read_file("binary.txt", workspace_root=tmp_path)

    assert _error_code(response) == "DECODE_ERROR"


def test_read_docx_extracts_paragraphs_and_tables(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Taylor Morgan")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Python"
    document.save(tmp_path / "resume.docx")

    response = read_file("resume.docx", workspace_root=tmp_path)
    data = _data(response)

    assert "Taylor Morgan" in data["content"]
    assert "Skill\tPython" in data["content"]
    assert data["extraction"]["table_count"] == 1


def test_read_pdf_extracts_text(tmp_path: Path) -> None:
    _write_minimal_text_pdf(tmp_path / "resume.pdf", "Python engineer")

    response = read_file("resume.pdf", workspace_root=tmp_path)
    data = _data(response)

    assert "Python engineer" in data["content"]
    assert data["extraction"]["page_count"] == 1


def test_read_corrupt_pdf_returns_parse_error(tmp_path: Path) -> None:
    (tmp_path / "broken.pdf").write_bytes(b"not a PDF")

    response = read_file("broken.pdf", workspace_root=tmp_path)

    assert _error_code(response) == "PARSE_ERROR"


def test_read_rejects_unsupported_extension(tmp_path: Path) -> None:
    (tmp_path / "resume.csv").write_text("Python", encoding="utf-8")

    response = read_file("resume.csv", workspace_root=tmp_path)

    assert _error_code(response) == "UNSUPPORTED_FILE_TYPE"


def test_read_enforces_file_size_limit(tmp_path: Path) -> None:
    (tmp_path / "resume.txt").write_text("12345", encoding="utf-8")

    response = read_file("resume.txt", workspace_root=tmp_path, max_file_size_bytes=4)

    assert _error_code(response) == "FILE_TOO_LARGE"


def test_read_rejects_invalid_file_size_limit(tmp_path: Path) -> None:
    (tmp_path / "resume.txt").write_text("content", encoding="utf-8")

    response = read_file("resume.txt", workspace_root=tmp_path, max_file_size_bytes=0)

    assert _error_code(response) == "INVALID_ARGUMENT"


def test_list_files_filters_and_sorts_case_insensitively(tmp_path: Path) -> None:
    (tmp_path / "zeta.TXT").write_text("z", encoding="utf-8")
    (tmp_path / "Alpha.txt").write_text("a", encoding="utf-8")
    (tmp_path / "resume.pdf").write_bytes(b"pdf")
    (tmp_path / ".hidden.txt").write_text("hidden", encoding="utf-8")
    (tmp_path / "nested").mkdir()

    response = list_files(".", ".txt", workspace_root=tmp_path)
    data = _data(response)

    assert data["extension"] == ".txt"
    assert data["count"] == 2
    assert [item["name"] for item in data["files"]] == ["Alpha.txt", "zeta.TXT"]


def test_list_files_returns_empty_collection_for_no_matches(tmp_path: Path) -> None:
    response = list_files(".", "pdf", workspace_root=tmp_path)
    data = _data(response)

    assert data["count"] == 0
    assert data["files"] == []


def test_write_file_creates_parent_directories_and_utf8_content(tmp_path: Path) -> None:
    response = write_file(
        "summaries/candidate.md",
        "Résumé summary",
        workspace_root=tmp_path,
    )
    data = _data(response)

    destination = tmp_path / "summaries" / "candidate.md"
    assert destination.read_text(encoding="utf-8") == "Résumé summary"
    assert data["bytes_written"] == len("Résumé summary".encode())
    assert data["overwritten"] is False
    assert list(destination.parent.glob(f".{destination.name}.*.tmp")) == []


def test_write_file_requires_explicit_overwrite(tmp_path: Path) -> None:
    destination = tmp_path / "summary.txt"
    destination.write_text("original", encoding="utf-8")

    rejected = write_file("summary.txt", "new", workspace_root=tmp_path)
    replaced = write_file("summary.txt", "new", workspace_root=tmp_path, overwrite=True)

    assert _error_code(rejected) == "FILE_ALREADY_EXISTS"
    assert _data(replaced)["overwritten"] is True
    assert destination.read_text(encoding="utf-8") == "new"


def test_write_file_rejects_path_outside_workspace(tmp_path: Path) -> None:
    response = write_file("../outside.txt", "blocked", workspace_root=tmp_path)

    assert _error_code(response) == "PATH_OUTSIDE_WORKSPACE"


def test_search_is_literal_case_insensitive_and_returns_context(tmp_path: Path) -> None:
    content = "First line\nExperienced with PYTHON. Uses Python daily."
    (tmp_path / "resume.txt").write_text(content, encoding="utf-8")

    response = search_in_file(
        "resume.txt",
        "python",
        workspace_root=tmp_path,
        context_characters=8,
    )
    data = _data(response)

    assert data["match_count"] == 2
    assert data["returned_match_count"] == 2
    assert [match["match"] for match in data["matches"]] == ["PYTHON", "Python"]
    assert data["matches"][0]["line_number"] == 2
    assert all("context" in match for match in data["matches"])


def test_search_treats_regex_characters_literally(tmp_path: Path) -> None:
    (tmp_path / "resume.txt").write_text("Uses C++ and Python", encoding="utf-8")

    response = search_in_file("resume.txt", "C++", workspace_root=tmp_path)

    assert _data(response)["match_count"] == 1


def test_search_returns_success_for_no_matches(tmp_path: Path) -> None:
    (tmp_path / "resume.txt").write_text("Java developer", encoding="utf-8")

    response = search_in_file("resume.txt", "Python", workspace_root=tmp_path)
    data = _data(response)

    assert data["match_count"] == 0
    assert data["matches"] == []


def test_search_limits_returned_matches_and_reports_total(tmp_path: Path) -> None:
    (tmp_path / "resume.txt").write_text("Python Python Python", encoding="utf-8")

    response = search_in_file(
        "resume.txt",
        "Python",
        workspace_root=tmp_path,
        max_matches=2,
    )
    data = _data(response)

    assert data["match_count"] == 3
    assert data["returned_match_count"] == 2
    assert response["warnings"] == ["Results were limited to 2 of 3 matches."]


def test_search_rejects_empty_keyword(tmp_path: Path) -> None:
    (tmp_path / "resume.txt").write_text("Python", encoding="utf-8")

    response = search_in_file("resume.txt", " ", workspace_root=tmp_path)

    assert _error_code(response) == "INVALID_KEYWORD"

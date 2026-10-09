"""Generate the fictional resume fixtures used by tests and demonstrations."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor, black
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


@dataclass(frozen=True, slots=True)
class Resume:
    filename: str
    name: str
    role: str
    contact: tuple[str, ...]
    summary: str
    skills: tuple[str, ...]
    experience: tuple[tuple[str, str, tuple[str, ...]], ...]
    education: tuple[str, ...]


RESUMES = (
    Resume(
        filename="aisha-khan-data-engineer.txt",
        name="Aisha Khan",
        role="Data Engineer",
        contact=("aisha.khan@example.com", "+1 202 555 0101", "Austin, Texas"),
        summary=(
            "Data engineer with five years of experience building reliable batch and streaming "
            "pipelines for analytics teams."
        ),
        skills=("Python", "SQL", "Apache Spark", "AWS", "Airflow", "Data modeling"),
        experience=(
            (
                "Data Engineer",
                "Northstar Analytics | 2022 to Present",
                (
                    "Built Python and Spark pipelines processing 40 million records per day.",
                    "Reduced failed Airflow jobs by adding validation and retry policies.",
                ),
            ),
            (
                "Junior Data Engineer",
                "Cedar Systems | 2020 to 2022",
                ("Maintained SQL transformations and AWS data-lake jobs.",),
            ),
        ),
        education=("Bachelor of Science in Computer Science, Lakeside University, 2020",),
    ),
    Resume(
        filename="liam-chen-frontend-developer.txt",
        name="Liam Chen",
        role="Frontend Developer",
        contact=("liam.chen@example.com", "+1 202 555 0102", "Seattle, Washington"),
        summary=(
            "Frontend developer focused on accessible design systems and responsive web "
            "applications."
        ),
        skills=("TypeScript", "React", "Next.js", "CSS", "Jest", "WCAG"),
        experience=(
            (
                "Frontend Developer",
                "Harbor Digital | 2021 to Present",
                (
                    "Created reusable React components used across four customer portals.",
                    "Added automated accessibility checks to the delivery pipeline.",
                ),
            ),
        ),
        education=("Bachelor of Design in Interaction Design, Western College, 2021",),
    ),
    Resume(
        filename="sofia-martinez-project-coordinator.txt",
        name="Sofia Martinez",
        role="Project Coordinator",
        contact=("sofia.martinez@example.com", "Denver, Colorado"),
        summary=(
            "Project coordinator experienced in software delivery schedules, risk logs, and "
            "stakeholder reporting."
        ),
        skills=("Jira", "Microsoft Excel", "Risk tracking", "Agile delivery", "Documentation"),
        experience=(
            (
                "Project Coordinator",
                "Juniper Works | 2023 to Present",
                (
                    "Coordinates weekly delivery plans for three product teams.",
                    "Maintains project risks, decisions, and release documentation.",
                ),
            ),
        ),
        education=(),
    ),
    Resume(
        filename="noah-williams-devops-engineer.docx",
        name="Noah Williams",
        role="DevOps Engineer",
        contact=("noah.williams@example.com", "+1 202 555 0104", "Raleigh, North Carolina"),
        summary=(
            "DevOps engineer with six years of experience operating cloud platforms and "
            "automating deployment workflows."
        ),
        skills=("PYTHON", "Docker", "Kubernetes", "Terraform", "Azure", "GitHub Actions"),
        experience=(
            (
                "DevOps Engineer",
                "Blue Peak Software | 2020 to Present",
                (
                    "Automated release checks with PYTHON and GitHub Actions.",
                    "Operated Kubernetes services and Terraform infrastructure in Azure.",
                ),
            ),
        ),
        education=("Bachelor of Science in Information Systems, Pine State University, 2020",),
    ),
    Resume(
        filename="priya-sharma-business-analyst.docx",
        name="Priya Sharma",
        role="Business Analyst",
        contact=("priya.sharma@example.com", "Chicago, Illinois"),
        summary=(
            "Business analyst who translates operational needs into measurable reporting and "
            "process requirements."
        ),
        skills=("SQL", "Power BI", "Microsoft Excel", "Requirements analysis", "Process mapping"),
        experience=(
            (
                "Business Analyst",
                "Fieldstone Services | 2022 to Present",
                (
                    "Defined reporting requirements with finance and operations teams.",
                    "Built Power BI dashboards backed by validated SQL queries.",
                ),
            ),
        ),
        education=("Master of Business Analytics, Metro University, 2022",),
    ),
    Resume(
        filename="ethan-brooks-machine-learning-engineer.pdf",
        name="Ethan Brooks",
        role="Machine Learning Engineer",
        contact=("ethan.brooks@example.com", "+1 202 555 0106", "Boston, Massachusetts"),
        summary=(
            "Machine learning engineer with four years of experience deploying forecasting and "
            "classification services."
        ),
        skills=("Python", "PyTorch", "scikit-learn", "FastAPI", "MLflow", "Docker"),
        experience=(
            (
                "Machine Learning Engineer",
                "Orchard Labs | 2022 to Present",
                (
                    "Deployed Python inference services with FastAPI and Docker.",
                    "Tracked experiments and model releases with MLflow.",
                ),
            ),
        ),
        education=("Master of Science in Data Science, Franklin Institute, 2022",),
    ),
    Resume(
        filename="maya-patel-java-engineer.pdf",
        name="Maya Patel",
        role="Java Software Engineer",
        contact=("maya.patel@example.com", "+1 202 555 0107", "Columbus, Ohio"),
        summary=(
            "Backend engineer with seven years of experience developing transaction services and "
            "integration APIs."
        ),
        skills=("Java", "Spring Boot", "PostgreSQL", "Kafka", "JUnit", "Basic python scripting"),
        experience=(
            (
                "Senior Software Engineer",
                "Summit Commerce | 2019 to Present",
                (
                    "Developed Spring Boot services handling payment and order events.",
                    "Used basic python scripts for migration checks and test-data preparation.",
                ),
            ),
        ),
        education=("Bachelor of Engineering in Software Engineering, Central Tech, 2019",),
    ),
    Resume(
        filename="olivia-reed-clinical-nurse.txt",
        name="Olivia Reed",
        role="Clinical Nurse",
        contact=("olivia.reed@example.com", "+1 202 555 0108", "Portland, Oregon"),
        summary=(
            "Registered nurse with six years of experience coordinating patient care across "
            "medical-surgical and outpatient settings."
        ),
        skills=("Patient assessment", "Care planning", "Epic", "Medication safety", "BLS"),
        experience=(
            (
                "Clinical Nurse",
                "Rosewood Health | 2020 to Present",
                (
                    "Coordinates daily care plans for a 24-bed medical-surgical unit.",
                    "Improved discharge handoffs through a standardized bedside checklist.",
                ),
            ),
        ),
        education=("Bachelor of Science in Nursing, Cascadia University, 2020",),
    ),
    Resume(
        filename="marcus-green-financial-analyst.txt",
        name="Marcus Green",
        role="Financial Analyst",
        contact=("marcus.green@example.com", "+1 202 555 0109", "Charlotte, North Carolina"),
        summary=(
            "Financial analyst with five years of experience in forecasting, variance analysis, "
            "and executive reporting."
        ),
        skills=("Financial modeling", "Microsoft Excel", "Power BI", "Budgeting", "SAP"),
        experience=(
            (
                "Senior Financial Analyst",
                "Granite Capital Services | 2021 to Present",
                (
                    "Owns quarterly forecasts for a 75 million dollar operating portfolio.",
                    "Automated monthly variance reporting and reduced preparation time by "
                    "30 percent.",
                ),
            ),
        ),
        education=("Bachelor of Science in Finance, Piedmont College, 2021",),
    ),
    Resume(
        filename="emma-wilson-digital-marketing-manager.txt",
        name="Emma Wilson",
        role="Digital Marketing Manager",
        contact=("emma.wilson@example.com", "+1 202 555 0110", "Nashville, Tennessee"),
        summary=(
            "Digital marketing manager specializing in lifecycle campaigns, content strategy, "
            "and measurable customer acquisition."
        ),
        skills=("Google Analytics", "SEO", "HubSpot", "Email marketing", "Campaign strategy"),
        experience=(
            (
                "Digital Marketing Manager",
                "Brightline Media | 2022 to Present",
                (
                    "Increased qualified organic traffic by 38 percent through a topic-cluster "
                    "strategy.",
                    "Built lifecycle campaigns that improved trial-to-paid conversion by "
                    "12 percent.",
                ),
            ),
        ),
        education=("Bachelor of Arts in Marketing, Cumberland University, 2022",),
    ),
    Resume(
        filename="aarav-mehta-hospitality-operations-manager.txt",
        name="Aarav Mehta",
        role="Hospitality Operations Manager",
        contact=("aarav.mehta@example.com", "+1 202 555 0111", "Orlando, Florida"),
        summary=(
            "Hospitality leader with eight years of experience improving guest service, staffing, "
            "and property operations."
        ),
        skills=(
            "Guest experience",
            "Revenue management",
            "Workforce planning",
            "Opera PMS",
            "Training",
        ),
        experience=(
            (
                "Operations Manager",
                "Suncrest Hotels | 2019 to Present",
                (
                    "Leads front-office and housekeeping operations for a 210-room property.",
                    "Raised guest satisfaction scores by nine points through service coaching.",
                ),
            ),
        ),
        education=("Bachelor of Hospitality Management, Atlantic College, 2018",),
    ),
    Resume(
        filename="zoe-turner-cybersecurity-analyst.docx",
        name="Zoe Turner",
        role="Cybersecurity Analyst",
        contact=("zoe.turner@example.com", "+1 202 555 0112", "Arlington, Virginia"),
        summary=(
            "Cybersecurity analyst with four years of experience in detection engineering, "
            "incident response, and cloud workload protection."
        ),
        skills=("SIEM", "Incident response", "Kubernetes", "AWS", "Vulnerability management"),
        experience=(
            (
                "Cybersecurity Analyst",
                "Sentinel Ridge | 2022 to Present",
                (
                    "Built SIEM detections for identity and cloud-control-plane events.",
                    "Reviewed Kubernetes workload policies and remediated exposed service "
                    "accounts.",
                ),
            ),
        ),
        education=("Bachelor of Science in Cybersecurity, Potomac University, 2022",),
    ),
    Resume(
        filename="daniel-kim-secondary-teacher.docx",
        name="Daniel Kim",
        role="Secondary Mathematics Teacher",
        contact=("daniel.kim@example.com", "+1 202 555 0113", "Sacramento, California"),
        summary=(
            "Secondary mathematics teacher with seven years of experience designing inclusive, "
            "standards-aligned instruction."
        ),
        skills=(
            "Curriculum design",
            "Classroom management",
            "Assessment",
            "Google Classroom",
            "Tutoring",
        ),
        experience=(
            (
                "Mathematics Teacher",
                "Riverbend Public Schools | 2019 to Present",
                (
                    "Teaches algebra and geometry to 140 students across five sections.",
                    "Introduced mastery-based review cycles that increased course pass rates.",
                ),
            ),
        ),
        education=("Master of Education, Capital State University, 2019",),
    ),
    Resume(
        filename="leila-hassan-ux-designer.docx",
        name="Leila Hassan",
        role="UX Designer",
        contact=("leila.hassan@example.com", "+1 202 555 0114", "Brooklyn, New York"),
        summary=(
            "UX designer with five years of experience turning customer research into accessible "
            "web and mobile product experiences."
        ),
        skills=("Figma", "User research", "Prototyping", "Design systems", "Accessibility"),
        experience=(
            (
                "UX Designer",
                "North Loop Products | 2021 to Present",
                (
                    "Led discovery and prototyping for a self-service onboarding redesign.",
                    "Expanded the design system with accessible form and navigation patterns.",
                ),
            ),
        ),
        education=("Bachelor of Fine Arts in Interaction Design, Prattview College, 2021",),
    ),
    Resume(
        filename="victor-alvarez-compliance-specialist.pdf",
        name="Victor Alvarez",
        role="Regulatory Compliance Specialist",
        contact=("victor.alvarez@example.com", "+1 202 555 0115", "Phoenix, Arizona"),
        summary=(
            "Compliance specialist with six years of experience managing controls, audits, and "
            "regulatory change programs."
        ),
        skills=(
            "Control testing",
            "Risk assessment",
            "Audit coordination",
            "Policy management",
            "GRC",
        ),
        experience=(
            (
                "Compliance Specialist",
                "Canyon Trust | 2020 to Present",
                (
                    "Coordinates annual control testing across lending and payment operations.",
                    "Maintains regulatory change assessments and remediation evidence in the "
                    "GRC platform.",
                ),
            ),
        ),
        education=("Bachelor of Business Administration, Desert State University, 2020",),
    ),
    Resume(
        filename="grace-okafor-supply-chain-manager.pdf",
        name="Grace Okafor",
        role="Supply Chain Manager",
        contact=("grace.okafor@example.com", "+1 202 555 0116", "Memphis, Tennessee"),
        summary=(
            "Supply chain manager with nine years of experience in demand planning, supplier "
            "performance, and distribution operations."
        ),
        skills=("Demand planning", "SAP", "Supplier management", "Inventory optimization", "Lean"),
        experience=(
            (
                "Supply Chain Manager",
                "Meridian Distribution | 2018 to Present",
                (
                    "Manages demand planning and replenishment for 1,200 active products.",
                    "Reduced backorders by 18 percent through supplier scorecards and "
                    "safety-stock reviews.",
                ),
            ),
        ),
        education=("Master of Science in Supply Chain Management, Delta University, 2018",),
    ),
    Resume(
        filename="lucas-morgan-renewable-energy-engineer.pdf",
        name="Lucas Morgan",
        role="Renewable Energy Engineer",
        contact=("lucas.morgan@example.com", "+1 202 555 0117", "Boulder, Colorado"),
        summary=(
            "Renewable energy engineer with five years of experience evaluating solar and battery "
            "projects from feasibility through commissioning."
        ),
        skills=("Solar design", "Battery storage", "AutoCAD", "PVsyst", "Energy modeling"),
        experience=(
            (
                "Renewable Energy Engineer",
                "High Plains Power | 2021 to Present",
                (
                    "Designed utility-scale solar layouts and energy-yield models for project "
                    "bids.",
                    "Supported commissioning tests for 40 megawatt-hours of battery storage.",
                ),
            ),
        ),
        education=("Bachelor of Science in Electrical Engineering, Front Range University, 2021",),
    ),
)

ADDITIONAL_RESUME_FILENAMES = frozenset(
    {
        "olivia-reed-clinical-nurse.txt",
        "marcus-green-financial-analyst.txt",
        "emma-wilson-digital-marketing-manager.txt",
        "aarav-mehta-hospitality-operations-manager.txt",
        "zoe-turner-cybersecurity-analyst.docx",
        "daniel-kim-secondary-teacher.docx",
        "leila-hassan-ux-designer.docx",
        "victor-alvarez-compliance-specialist.pdf",
        "grace-okafor-supply-chain-manager.pdf",
        "lucas-morgan-renewable-energy-engineer.pdf",
    }
)


def generate_all(output_directory: Path) -> list[Path]:
    """Generate every sample resume and return its path."""

    return _generate(RESUMES, output_directory)


def generate_additional(output_directory: Path) -> list[Path]:
    """Generate only the ten cross-domain resumes added after the starter set."""

    resumes = tuple(resume for resume in RESUMES if resume.filename in ADDITIONAL_RESUME_FILENAMES)
    return _generate(resumes, output_directory)


def _generate(resumes: tuple[Resume, ...], output_directory: Path) -> list[Path]:
    """Generate a selected collection of resume fixtures."""

    output_directory.mkdir(parents=True, exist_ok=True)
    generated: list[Path] = []
    for resume in resumes:
        destination = output_directory / resume.filename
        extension = destination.suffix.lower()
        if extension == ".txt":
            _write_txt(resume, destination)
        elif extension == ".docx":
            _write_docx(resume, destination)
        elif extension == ".pdf":
            _write_pdf(resume, destination)
        else:
            raise ValueError(f"Unsupported output format: {extension}")
        generated.append(destination)
    return generated


def _write_txt(resume: Resume, destination: Path) -> None:
    lines = [resume.name, resume.role, *resume.contact, "", "PROFESSIONAL SUMMARY", resume.summary]
    lines.extend(("", "SKILLS", ", ".join(resume.skills), "", "EXPERIENCE"))
    for title, employer, bullets in resume.experience:
        lines.extend((title, employer))
        lines.extend(f"- {bullet}" for bullet in bullets)
    if resume.education:
        lines.extend(("", "EDUCATION", *resume.education))
    destination.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _set_run_font(run: object, name: str, size: float, *, bold: bool = False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _write_docx(resume: Resume, destination: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
    for style_name in ("Title", "Heading 1"):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    title_style_properties = styles["Title"].element.pPr
    if title_style_properties is not None:
        title_border = title_style_properties.find(qn("w:pBdr"))
        if title_border is not None:
            title_style_properties.remove(title_border)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    _set_run_font(title.add_run(resume.name), "Arial", 20, bold=True)

    role = document.add_paragraph()
    role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    role.paragraph_format.space_after = Pt(4)
    _set_run_font(role.add_run(resume.role), "Arial", 12, bold=True)

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(10)
    _set_run_font(contact.add_run(" | ".join(resume.contact)), "Arial", 9.5)

    _add_docx_section(document, "Professional Summary")
    paragraph = document.add_paragraph(resume.summary)
    paragraph.paragraph_format.space_after = Pt(7)

    _add_docx_section(document, "Skills")
    skills = document.add_paragraph(", ".join(resume.skills))
    skills.paragraph_format.space_after = Pt(7)

    _add_docx_section(document, "Experience")
    for title_text, employer, bullets in resume.experience:
        heading = document.add_paragraph()
        heading.paragraph_format.space_after = Pt(1)
        _set_run_font(heading.add_run(title_text), "Arial", 10.5, bold=True)
        employer_line = document.add_paragraph(employer)
        employer_line.paragraph_format.space_after = Pt(2)
        for bullet in bullets:
            item = document.add_paragraph(bullet, style="List Bullet")
            item.paragraph_format.space_after = Pt(2)

    if resume.education:
        _add_docx_section(document, "Education")
        for item in resume.education:
            document.add_paragraph(item)

    document.core_properties.author = "LLM File System Assistant Sample Data"
    document.core_properties.title = f"{resume.name} Resume"
    document.save(destination)


def _add_docx_section(document: Document, title: str) -> None:
    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.space_before = Pt(7)
    heading.paragraph_format.space_after = Pt(3)
    _set_run_font(heading.add_run(title), "Arial", 12, bold=True)


def _write_pdf(resume: Resume, destination: Path) -> None:
    document = SimpleDocTemplate(
        str(destination),
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=f"{resume.name} Resume",
        author="LLM File System Assistant Sample Data",
    )
    sample_styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ResumeTitle",
        parent=sample_styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=23,
        textColor=black,
        alignment=TA_CENTER,
        spaceAfter=8,
    )
    role_style = ParagraphStyle(
        "ResumeRole",
        parent=sample_styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        alignment=TA_CENTER,
        spaceAfter=5,
    )
    contact_style = ParagraphStyle(
        "ResumeContact",
        parent=sample_styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        alignment=TA_CENTER,
        spaceAfter=10,
    )
    heading_style = ParagraphStyle(
        "ResumeHeading",
        parent=sample_styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=HexColor("#000000"),
        spaceBefore=7,
        spaceAfter=3,
    )
    body_style = ParagraphStyle(
        "ResumeBody",
        parent=sample_styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=black,
        spaceAfter=4,
    )
    bullet_style = ParagraphStyle(
        "ResumeBullet",
        parent=body_style,
        leftIndent=14,
        firstLineIndent=-7,
        bulletIndent=0,
        spaceAfter=2,
    )

    story: list[object] = [
        Paragraph(resume.name, title_style),
        Paragraph(resume.role, role_style),
        Paragraph(" | ".join(resume.contact), contact_style),
        Paragraph("Professional Summary", heading_style),
        Paragraph(resume.summary, body_style),
        Paragraph("Skills", heading_style),
        Paragraph(", ".join(resume.skills), body_style),
        Paragraph("Experience", heading_style),
    ]
    for title_text, employer, bullets in resume.experience:
        story.extend(
            (
                Paragraph(f"<b>{title_text}</b>", body_style),
                Paragraph(employer, body_style),
            )
        )
        story.extend(Paragraph(bullet, bullet_style, bulletText="-") for bullet in bullets)
    if resume.education:
        story.extend((Paragraph("Education", heading_style),))
        story.extend(Paragraph(item, body_style) for item in resume.education)
    story.append(Spacer(1, 2))
    document.build(story)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("data/resumes"),
        help="Directory that receives the generated sample resumes.",
    )
    arguments = parser.parse_args()
    generated = generate_all(arguments.output_dir.resolve())
    for path in generated:
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
